import asyncio
import io
import sys
import threading
import time
import zipfile
from pathlib import Path
from types import SimpleNamespace

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from fastapi import UploadFile
from routes import upload


def test_general_text_attachment_preserves_content_and_reports_truncation():
    result = asyncio.run(upload.general_agent_upload(
        UploadFile(filename="../notes.md", file=io.BytesIO(("论证甲" * 8000).encode())), _g={}))
    assert result["filename"] == "notes.md"
    assert result["content"] == ("论证甲" * 8000)[:20000]
    assert result["truncated"] is True
    assert result["extracted_chars"] == 24000


def test_empty_unsupported_and_oversized_are_errors():
    for filename, data, code in [("empty.txt", b"", 400), ("run.exe", b"binary", 415),
                                 ("large.txt", b"x" * (20 * 1024 * 1024 + 1), 413)]:
        result = asyncio.run(upload.general_agent_upload(
            UploadFile(filename=filename, file=io.BytesIO(data)), _g={}))
        assert result.status_code == code


def test_conversion_does_not_block_event_loop(monkeypatch):
    entered, finish = threading.Event(), threading.Event()
    def convert(*args):
        entered.set()
        assert finish.wait(2)
        return {"content": "extracted", "filename": "a.pdf"}
    monkeypatch.setattr(upload, "_convert_general_attachment", convert)
    async def run():
        task = asyncio.create_task(upload.general_agent_upload(
            UploadFile(filename="a.pdf", file=io.BytesIO(b"pdf")), _g={}))
        for _ in range(100):
            if entered.is_set():
                break
            await asyncio.sleep(0.005)
        assert entered.is_set() and not task.done()
        finish.set()
        return await task
    assert asyncio.run(run())["content"] == "extracted"


def test_same_named_concurrent_documents_use_isolated_cleaned_temp_files(monkeypatch):
    barrier = threading.Barrier(2)
    paths = []
    class Converter:
        def convert(self, filename):
            path = Path(filename)
            paths.append(path)
            barrier.wait(timeout=2)
            return SimpleNamespace(text_content=path.read_text())
    monkeypatch.setitem(sys.modules, "markitdown", SimpleNamespace(MarkItDown=Converter))
    async def run():
        return await asyncio.gather(*[
            asyncio.to_thread(upload._convert_general_attachment, text, "same.docx", ".docx")
            for text in (b"first", b"second")])
    result = asyncio.run(run())
    assert [item["content"] for item in result] == ["first", "second"]
    assert len(set(paths)) == 2
    assert all(not path.exists() for path in paths)


def test_real_docx_converter_dependencies_are_installed():
    data = io.BytesIO()
    with zipfile.ZipFile(data, "w") as doc:
        doc.writestr("[Content_Types].xml", '''<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"><Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/><Default Extension="xml" ContentType="application/xml"/><Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/></Types>''')
        doc.writestr("_rels/.rels", '''<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/></Relationships>''')
        doc.writestr("word/document.xml", '''<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body><w:p><w:r><w:t>真实附件：事实与解释不能互相替代。</w:t></w:r></w:p></w:body></w:document>''')
    result = upload._convert_general_attachment(data.getvalue(), "example.docx", ".docx")
    assert "事实与解释不能互相替代" in result["content"]


def test_image_data_uses_its_actual_format_and_preserves_bytes(monkeypatch):
    from PIL import Image
    buf = io.BytesIO()
    Image.new("RGB", (3, 3), "white").save(buf, format="PNG")
    raw = buf.getvalue()
    seen = []
    monkeypatch.setattr(upload, "_agnes_vision", lambda data, **kwargs: seen.append((data, kwargs)) or "识别内容")
    result = upload._convert_general_attachment(raw, "mislabeled.jpg", ".jpg")
    assert seen == [(raw, {"mime_type": "image/png", "prompt": upload._GENERAL_VISION_PROMPT})]
    assert result["kind"] == "image" and result["content"] == "识别内容"
    import pytest
    with pytest.raises(ValueError, match="图片内容"):
        upload._convert_general_attachment(b"not an image", "bad.png", ".png")


def test_real_powerpoint_and_spreadsheet_converters():
    from pptx import Presentation
    from openpyxl import Workbook
    presentation = Presentation()
    presentation.slides.add_slide(presentation.slide_layouts[5]).shapes.title.text = "Argument test"
    buf = io.BytesIO()
    presentation.save(buf)
    assert "Argument test" in upload._convert_general_attachment(buf.getvalue(), "test.pptx", ".pptx")["content"]
    workbook = Workbook()
    workbook.active.append(["论点", "限制"])
    workbook.active.append(["事实", "仍需解释"])
    buf = io.BytesIO()
    workbook.save(buf)
    assert "仍需解释" in upload._convert_general_attachment(buf.getvalue(), "test.xlsx", ".xlsx")["content"]


def test_real_pdf_text_extraction():
    stream = b"BT /F1 12 Tf 50 750 Td (Actual PDF extraction fixture.) Tj ET"
    objects = [b"<< /Type /Catalog /Pages 2 0 R >>", b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
               b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>",
               b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
               b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream"]
    pdf, offsets = bytearray(b"%PDF-1.4\n"), []
    for index, obj in enumerate(objects, 1):
        offsets.append(len(pdf))
        pdf.extend(str(index).encode() + b" 0 obj\n" + obj + b"\nendobj\n")
    xref = len(pdf)
    pdf.extend(b"xref\n0 6\n0000000000 65535 f \n")
    for offset in offsets:
        pdf.extend(f"{offset:010} 00000 n \n".encode())
    pdf.extend(f"trailer\n<< /Size 6 /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF".encode())
    assert "Actual PDF extraction fixture" in upload._convert_general_attachment(bytes(pdf), "test.pdf", ".pdf")["content"]
